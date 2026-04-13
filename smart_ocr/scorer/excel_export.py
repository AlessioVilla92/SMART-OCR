"""
scorer/excel_export.py

Popola il template Excel CBCL_6-18.xlt con le risposte OCR.
Le formule gia' presenti nel template calcolano automaticamente
i punteggi sindromici, broadband e DSM.
"""

import shutil
from pathlib import Path
from openpyxl import load_workbook
from .cbcl_scorer import CBCLScorer, Compilatore, item_to_excel_row, ALL_ITEMS


def export_to_excel(
    responses: dict,
    compilatore: Compilatore,
    template_path: str = "templates/CBCL_6-18.xlt",
    output_path: str = None,
    child_name: str = "",
    child_dob: str = "",
    test_date: str = "",
) -> str:
    """
    Copia il template e lo popola con le risposte OCR.

    Args:
        responses: {item: 0/1/2} -- risposte OCR
        compilatore: MADRE o PADRE
        template_path: percorso del template .xlt
        output_path: percorso output (auto-generato se None)
        child_name: nome del bambino (opzionale, per cella E3)
        child_dob: data nascita (opzionale, per cella E4)
        test_date: data test (opzionale, per cella E5)

    Returns:
        Percorso del file Excel generato.
    """
    from datetime import datetime

    template = Path(template_path)
    if not template.exists():
        raise FileNotFoundError(f"Template non trovato: {template_path}")

    if output_path is None:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        comp = compilatore.value
        output_path = f"output/CBCL_{comp}_{ts}.xlsx"

    # Copia template
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, out)

    # Apri e popola
    wb = load_workbook(str(out))
    ws = wb["Foglio1"]

    # Colonna dati: B per Madre, C per Padre
    col_idx = 2 if compilatore == Compilatore.MADRE else 3

    # Scrivi risposte
    for item in ALL_ITEMS:
        val = responses.get(item)
        if val is not None:
            row = item_to_excel_row(item)
            ws.cell(row=row, column=col_idx, value=int(val))

    # Scrivi metadati (opzionali)
    if child_name:
        ws["E3"] = child_name
    if child_dob:
        ws["E4"] = child_dob
    if test_date:
        ws["E5"] = test_date

    wb.save(str(out))
    return str(out)
